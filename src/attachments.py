"""
attachments.py — Ek dosyalardan metin cikarma.

Projenin en kritik teknik noktasi burasi: tedarikcilerin cogu fiyati mail
govdesine yazmaz, PDF/Excel eki atar. Bir kismi da taranmis PDF gonderir —
o dosyada metin katmani YOKTUR, pdftotext bos doner. Bunu sessizce gecmek
yerine acikca tespit edip "insan baksin" diye isaretliyoruz.
"""

import subprocess
from pathlib import Path

DESTEKLI = {".pdf", ".xlsx", ".xls", ".xlsm", ".csv", ".docx", ".txt"}


def _pdf(yol):
    try:
        r = subprocess.run(["pdftotext", "-layout", str(yol), "-"],
                           capture_output=True, text=True, timeout=90)
        metin = r.stdout or ""
    except Exception:
        metin = ""
    if len(metin.strip()) < 30:
        try:
            import pdfplumber
            parcalar = []
            with pdfplumber.open(str(yol)) as pdf:
                for sayfa in pdf.pages[:10]:
                    parcalar.append(sayfa.extract_text() or "")
                    for tablo in (sayfa.extract_tables() or []):
                        for satir in tablo:
                            parcalar.append(" | ".join(h or "" for h in satir))
            metin = "\n".join(parcalar)
        except Exception:
            pass
    return metin


def _excel(yol):
    import pandas as pd
    parcalar = []
    try:
        if Path(yol).suffix.lower() == ".csv":
            sayfalar = {"csv": pd.read_csv(yol, dtype=object, sep=None, engine="python")}
        else:
            sayfalar = pd.read_excel(yol, sheet_name=None, dtype=object)
        for ad, df in sayfalar.items():
            parcalar.append(f"--- {ad} ---")
            parcalar.append(df.dropna(how="all").to_string(index=False, max_rows=200))
    except Exception as e:
        parcalar.append(f"(okunamadi: {e})")
    return "\n".join(parcalar)


def _docx(yol):
    try:
        from docx import Document
        d = Document(str(yol))
        parcalar = [p.text for p in d.paragraphs]
        for t in d.tables:
            for satir in t.rows:
                parcalar.append(" | ".join(h.text for h in satir.cells))
        return "\n".join(parcalar)
    except Exception:
        return ""


def metin_cikar(yol):
    """
    Doner: (metin, uyari). uyari None degilse insan mudahalesi gerekiyor.
    """
    p = Path(yol)
    uzanti = p.suffix.lower()
    if uzanti not in DESTEKLI:
        return "", f"Desteklenmeyen ek turu: {p.name}"

    if uzanti == ".pdf":
        metin = _pdf(p)
    elif uzanti in (".xlsx", ".xls", ".xlsm", ".csv"):
        metin = _excel(p)
    elif uzanti == ".docx":
        metin = _docx(p)
    else:
        metin = p.read_text(encoding="utf-8", errors="ignore")

    if len(metin.strip()) < 30:
        return metin, (f"{p.name} icinden metin cikmadi — muhtemelen taranmis "
                       f"(goruntu) dosya. Elle bakilmasi gerekiyor.")
    return metin, None
